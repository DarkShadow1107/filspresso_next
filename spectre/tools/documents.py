import io
import json
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path

FILES_DIR = Path("/srv/tools/files")

FORMATS = {"txt", "md", "docx", "pdf", "pptx", "xlsx"}


def _slug(name: str) -> str:
    cleaned = re.sub(r"[^\w.-]+", "-", name.strip()).strip("-.")
    return cleaned[:80] or "document"


def _paragraphs(content: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n{2,}", content.strip()) if part.strip()]


def _rows(content: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in content.strip().splitlines():
        cells = [cell.strip() for cell in re.split(r"\t|,", line)]
        if any(cells):
            rows.append(cells)
    return rows or [["empty"]]


def _render_docx(title: str, content: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=0)
    for paragraph in _paragraphs(content):
        if paragraph.startswith("## "):
            document.add_heading(paragraph[3:], level=2)
        elif paragraph.startswith("# "):
            document.add_heading(paragraph[2:], level=1)
        elif paragraph.startswith("- "):
            document.add_paragraph(paragraph[2:], style="List Bullet")
        else:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_pdf(title: str, content: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    buffer = io.BytesIO()
    story: list = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for paragraph in _paragraphs(content):
        clean = paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(clean.replace("\n", "<br/>"), styles["BodyText"]))
        story.append(Spacer(1, 6))
    SimpleDocTemplate(buffer, pagesize=letter, title=title).build(story)
    return buffer.getvalue()


def _render_pptx(title: str, content: str) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slides = [slide.strip() for slide in content.split("\n---\n") if slide.strip()] or [" "]
    for index, slide_body in enumerate(slides):
        layout_index = 0 if index == 0 else 1
        slide = presentation.slides.add_slide(presentation.slide_layouts[layout_index])
        lines = [line.strip() for line in slide_body.splitlines() if line.strip()]
        slide.shapes.title.text = (lines[0].lstrip("# ") if lines else title)[:200]
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = lines[1:] or lines
        for bullet_index, bullet in enumerate(bullets[:12]):
            paragraph = body.paragraphs[0] if bullet_index == 0 else body.add_paragraph()
            paragraph.text = bullet.lstrip("- ").strip()[:300]
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _render_xlsx(title: str, content: str) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (_slug(title) or "Sheet1")[:31]
    for row in _rows(content):
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_document(fmt: str, title: str, content: str) -> bytes:
    fmt = fmt.lower().lstrip(".")
    if fmt in {"txt", "md"}:
        return content.encode("utf-8")
    if fmt == "docx":
        return _render_docx(title, content)
    if fmt == "pdf":
        return _render_pdf(title, content)
    if fmt == "pptx":
        return _render_pptx(title, content)
    if fmt == "xlsx":
        return _render_xlsx(title, content)
    raise ValueError(f"Unsupported format: {fmt}")


def read_registry() -> dict:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    path = FILES_DIR / "registry.json"
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def write_registry(registry: dict) -> None:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    (FILES_DIR / "registry.json").write_text(json.dumps(registry, indent=2), encoding="utf-8")


def save_document(filename: str, fmt: str, content: str, *, modified_from: str | None = None) -> dict:
    fmt = fmt.lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError(f"Unsupported format: {fmt}")
    title = Path(filename).stem or "Document"
    blob = render_document(fmt, title, content)
    document_id = str(uuid.uuid4())
    stored_name = f"{document_id}.{fmt}"
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    (FILES_DIR / stored_name).write_bytes(blob)
    registry = read_registry()
    entry = {
        "id": document_id,
        "filename": filename if filename.endswith(f".{fmt}") else f"{_slug(filename)}.{fmt}",
        "format": fmt,
        "title": title,
        "bytes": len(blob),
        "created_at": datetime.now(UTC).isoformat(),
        "modified_from": modified_from,
        "content_chars": len(content),
    }
    registry[document_id] = entry
    write_registry(registry)
    return entry


def modify_document(document_id: str, content: str) -> dict:
    previous = read_registry().get(document_id)
    if not previous:
        raise KeyError("Document not found")
    return save_document(previous["filename"], previous["format"], content, modified_from=document_id)


def get_entry(document_id: str) -> dict | None:
    return read_registry().get(document_id)


def read_bytes(entry: dict) -> bytes:
    return (FILES_DIR / f"{entry['id']}.{entry['format']}").read_bytes()


def list_documents() -> list[dict]:
    return sorted(read_registry().values(), key=lambda item: item["created_at"], reverse=True)
