import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models import DocumentChunk, Source, SourceVersion
from app.services.ingestion import chunk_text

MAX_IMAGE_BYTES = 8 * 1024 * 1024

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
HTML_SUFFIXES = {".html", ".htm"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
PPTX_SUFFIXES = {".pptx"}
XLSX_SUFFIXES = {".xlsx"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}

MIME_BY_SUFFIX = {
    **{suffix: "text/plain; charset=utf-8" for suffix in TEXT_SUFFIXES},
    **{suffix: "text/html" for suffix in HTML_SUFFIXES},
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tif": "image/tiff",
    ".tiff": "image/tiff",
}

_TAG_RE = re.compile(r"<(script|style|noscript|svg|iframe)\b.*?</\1>", re.IGNORECASE | re.DOTALL)
_HTML_RE = re.compile(r"<[^>]+>")


class UnsupportedDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class DocumentParse:
    kind: str
    mime_type: str
    text: str


def _extract_html(data: bytes) -> str:
    decoded = data.decode("utf-8", errors="replace")
    without_hidden = _TAG_RE.sub(" ", decoded)
    text = _HTML_RE.sub(" ", without_hidden)
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def _extract_pdf(data: bytes) -> str:
    import fitz

    parts: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        for page in document:
            parts.append(page.get_text("text"))
    return "\n".join(part for part in parts if part.strip())


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts.append(f"[Slide {index}]")
        for shape in slide.shapes:
            frame_text = getattr(shape, "text", "")
            if isinstance(frame_text, str) and frame_text.strip():
                parts.append(frame_text.strip())
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None and str(value).strip()]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts)


def parse_document(filename: str, data: bytes) -> DocumentParse:
    suffix = Path(filename).suffix.lower()
    if not suffix:
        raise UnsupportedDocumentError("File has no recognizable extension")
    if len(data) == 0:
        raise UnsupportedDocumentError("File is empty")

    if suffix in IMAGE_SUFFIXES:
        if len(data) > MAX_IMAGE_BYTES:
            raise UnsupportedDocumentError("Image exceeds the 8 MB sandbox analysis limit")
        return DocumentParse(kind="image", mime_type=MIME_BY_SUFFIX[suffix], text="")

    known = (
        TEXT_SUFFIXES | HTML_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES | PPTX_SUFFIXES | XLSX_SUFFIXES
    )
    if suffix not in known:
        supported = ", ".join(sorted(known | IMAGE_SUFFIXES))
        raise UnsupportedDocumentError(f"Unsupported file type. Supported: {supported}")

    settings_limit = get_settings().max_upload_bytes
    if len(data) > settings_limit:
        raise UnsupportedDocumentError("File exceeds the upload limit")

    if suffix in TEXT_SUFFIXES:
        text = data.decode("utf-8", errors="replace")
    elif suffix in HTML_SUFFIXES:
        text = _extract_html(data)
    elif suffix in PDF_SUFFIXES:
        text = _extract_pdf(data)
    elif suffix in DOCX_SUFFIXES:
        text = _extract_docx(data)
    elif suffix in PPTX_SUFFIXES:
        text = _extract_pptx(data)
    else:
        text = _extract_xlsx(data)

    if len(text.strip()) < 20:
        raise UnsupportedDocumentError(
            "No extractable text found; scanned PDFs are not OCR'd on upload"
        )
    return DocumentParse(kind="document", mime_type=MIME_BY_SUFFIX[suffix], text=text)


def _minio_endpoint_url() -> str:
    endpoint = get_settings().minio_endpoint.rstrip("/")
    if endpoint.startswith(("http://", "https://")):
        return endpoint
    return f"http://{endpoint}"


def document_store_options() -> dict[str, Any]:
    settings = get_settings()
    return {
        "endpoint_url": _minio_endpoint_url(),
        "aws_access_key_id": settings.minio_access_key,
        "aws_secret_access_key": settings.minio_secret_key,
        "region_name": "us-east-1",
    }


def safe_filename(filename: str) -> str:
    keep = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name).strip("._")
    return keep[:120] or "upload.bin"


async def put_document_original(key: str, data: bytes, content_type: str) -> None:
    from aiobotocore.session import get_session

    settings = get_settings()
    session = get_session()
    async with session.create_client("s3", **document_store_options()) as client:
        try:
            await client.head_bucket(Bucket=settings.minio_bucket)
        except Exception:
            await client.create_bucket(Bucket=settings.minio_bucket)
        await client.put_object(
            Bucket=settings.minio_bucket, Key=key, Body=data, ContentType=content_type
        )


async def delete_document_originals(keys: list[str]) -> None:
    from aiobotocore.session import get_session

    if not keys:
        return
    settings = get_settings()
    session = get_session()
    async with session.create_client("s3", **document_store_options()) as client:
        for key in keys:
            try:
                await client.delete_object(Bucket=settings.minio_bucket, Key=key)
            except Exception:
                continue


async def persist_user_document(
    db: AsyncSession,
    *,
    organization_id: UUID,
    filename: str,
    data: bytes,
) -> tuple[Source, SourceVersion]:
    parsed = parse_document(filename, data)
    import hashlib

    content_hash = hashlib.sha256(data).hexdigest()
    canonical_url = f"local://document/{uuid4()}"

    source = Source(
        organization_id=organization_id,
        canonical_url=canonical_url,
        title=Path(filename).name[:500],
        source_tier="user_provided",
    )
    db.add(source)
    await db.flush()

    storage_key: str | None = None
    if parsed.kind == "image":
        storage_key = f"documents/{organization_id}/{source.id}/{safe_filename(filename)}"
        await put_document_original(storage_key, data, parsed.mime_type)

    version = SourceVersion(
        source_id=source.id,
        content_hash=content_hash,
        extracted_text=parsed.text,
        metadata_json={
            "filename": Path(filename).name,
            "mime_type": parsed.mime_type,
            "kind": parsed.kind,
            "size_bytes": len(data),
            "storage_key": storage_key,
            "origin": "user_upload",
        },
    )
    db.add(version)
    await db.flush()

    if parsed.text:
        for ordinal, content in enumerate(chunk_text(parsed.text)):
            db.add(
                DocumentChunk(
                    organization_id=organization_id,
                    source_version_id=version.id,
                    ordinal=ordinal,
                    content=content,
                    metadata_json={
                        "chunking": "recursive-1000-160",
                        "source_url": canonical_url,
                    },
                )
            )
    return source, version
